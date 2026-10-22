"""XLS/XLSX real-byte contracts and atomic import, not a CSV replacement."""
from contextlib import closing
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import hashlib,json,sqlite3,tempfile,unittest,uuid,warnings,zipfile
from openpyxl import Workbook,load_workbook
import xlwt
from services.reference_errors import DemoError
from services.replica_risk_excel import HEADERS,FIELDS,read_rows,workbook_bytes
from services.replica_risk_units import ProjectDirectory,RiskPrincipal,RiskUnitStore

def key():return uuid.uuid4().hex
def fixture(rows,extension='xlsx',headers=HEADERS):
 if extension=='xlsx':
  w=Workbook();s=w.active;s.append(headers)
  for row in rows:s.append(row)
  out=BytesIO();w.save(out);w.close();return out.getvalue()
 w=xlwt.Workbook();s=w.add_sheet('sheet')
 for i,row in enumerate([headers,*rows]):
  for j,value in enumerate(row):
   if value is not None:s.write(i,j,value)
 out=BytesIO();w.save(out);return out.getvalue()

class ExcelTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.s=RiskUnitStore(Path(self.tmp.name)/'risk.sqlite')
  self.d=ProjectDirectory('p',[{'id':'d1','name':'合成一部'},{'id':'d2','name':'合成二部'}],[{'id':'tester','name':'合成甲','departments':['d1']},{'id':'u2','name':'合成乙','departments':['d2']}],[{'code':'1','name':'重大风险'},{'code':'4','name':'低风险'}],'dir-1');self.p=RiskPrincipal('p','tester',frozenset('risk-unit:'+x for x in ('read','create','update','delete','audit','export','import')),frozenset(['d1','d2']),'policy-1')
  self.row=['合成一部','合成甲','合成装置','合成单元','1','114.025,32.98']
 def error(self,status,fn,*args,**kwargs):
  with self.assertRaises(DemoError) as c:fn(*args,**kwargs)
  self.assertEqual(c.exception.status,status);return c.exception.message
 def import_bytes(self,raw,extension='xlsx',request=None,principal=None):return self.s.import_rows(principal or self.p,self.d,request or key(),read_rows('risk.'+extension,raw,self.d),hashlib.sha256(raw).hexdigest())
 def counts(self):
  with closing(self.s.connect()) as db:return [db.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in ('rp_risk_units','rp_risk_events','rp_risk_requests','rp_risk_outbox')]
 def test_original_six_column_empty_template(self):
  w=load_workbook(BytesIO(workbook_bytes()));self.assertEqual(w.sheetnames,['sheet']);self.assertEqual(tuple(c.value for c in w.active[1]),HEADERS);self.assertEqual(w.active.max_row,1);w.close()
 def test_xls_and_xlsx_import_real_bytes(self):
  for extension in ('xls','xlsx'):
   r=self.import_bytes(fixture([self.row],extension),extension);self.assertEqual(r['imported'],1)
  self.assertEqual(self.counts(),[2,2,2,2])
 def test_reordered_headers_blank_rows_and_dictionary_label(self):
  row=list(self.row);row[4]='重大风险';raw=fixture([row[::-1],[None]*6,row[::-1]],headers=HEADERS[::-1]);entries=read_rows('r.xlsx',raw,self.d);self.assertEqual([r['rowNumber'] for r in entries],[2,4]);self.assertEqual(entries[0]['values']['riskLevel'],'1');self.assertEqual(self.import_bytes(raw)['imported'],2)
 def test_batch_provenance_shared_and_each_row_audited(self):
  raw=fixture([self.row,self.row]);r=self.import_bytes(raw)
  for index,identity in enumerate(r['ids'],2):
   entry=self.s.history(self.p,self.d,identity)[0]['payload']['importBatch'];self.assertEqual(entry['id'],r['batchId']);self.assertEqual(entry['sourceSHA256'],hashlib.sha256(raw).hexdigest());self.assertEqual(entry['rowNumber'],index);self.assertEqual(entry['rowCount'],2)
 def test_exact_batch_replay_no_duplicates(self):
  raw=fixture([self.row,self.row]);request=key();first=self.import_bytes(raw,request=request);self.assertEqual(self.import_bytes(raw,request=request),first);self.assertEqual(self.counts(),[2,2,2,2])
 def test_changed_tail_or_count_same_key_rejected_atomically(self):
  raw=fixture([self.row,self.row]);request=key();self.import_bytes(raw,request=request)
  for rows in ([self.row],[self.row,self.row,self.row],[self.row,[*self.row[:3],'另一单元',*self.row[4:]]]):self.error(409,self.import_bytes,fixture(rows),request=request)
  self.assertEqual(self.counts(),[2,2,2,2])
 def test_bad_last_row_reports_line_and_leaves_no_partial_import(self):
  bad=list(self.row);bad[1]='合成乙';message=self.error(400,self.import_bytes,fixture([self.row,[None]*6,bad]));self.assertIn('第4行',message);self.assertEqual(self.counts(),[0,0,0,0])
 def test_mid_transaction_failure_rolls_back_prior_valid_rows(self):
  second=list(self.row);second[3]='第二行故障'
  with closing(self.s.connect()) as db:db.execute("CREATE TRIGGER import_failure BEFORE INSERT ON rp_risk_outbox WHEN json_extract(NEW.payload,'$.record.analysisUnit')='第二行故障' BEGIN SELECT RAISE(ABORT,'injected'); END")
  with self.assertRaises(sqlite3.IntegrityError):self.import_bytes(fixture([self.row,second]))
  self.assertEqual(self.counts(),[0,0,0,0])
 def test_import_permission_scope_and_replay_revocation(self):
  raw=fixture([self.row]);self.error(403,self.import_bytes,raw,principal=replace(self.p,permissions=self.p.permissions-{'risk-unit:import'}));self.error(403,self.import_bytes,raw,principal=replace(self.p,departments=frozenset(['d2'])));request=key();self.import_bytes(raw,request=request);self.error(403,self.import_bytes,raw,request=request,principal=replace(self.p,departments=frozenset(['d2'])));self.assertEqual(self.counts(),[1,1,1,1])
 def test_headers_unknown_duplicate_and_empty_data(self):
  for headers in ([*HEADERS[:-1],'未知列'],[*HEADERS[:-1],HEADERS[0]],HEADERS[:5]):self.error(400,read_rows,'r.xlsx',fixture([self.row],headers=headers),self.d)
  self.error(400,read_rows,'r.xlsx',fixture([]),self.d);self.assertEqual(self.counts(),[0,0,0,0])
 def test_magic_extension_filename_corrupt_and_missing_values(self):
  raw=fixture([self.row]);self.error(400,read_rows,'r.xls',raw,self.d);self.error(400,read_rows,'../r.xlsx',raw,self.d);self.error(400,read_rows,'r.xlsx',b'PK\x03\x04broken',self.d);self.error(415,read_rows,'r.csv',b'a,b',self.d)
  for index in range(6):
   bad=list(self.row);bad[index]=None;self.error(400,self.import_bytes,fixture([bad]))
 def test_first_sheet_and_text_literal_export_roundtrip(self):
  r=self.import_bytes(fixture([self.row]));record=self.s.get(self.p,self.d,r['ids'][0]);record['analysisObject']='=HYPERLINK("https://example.invalid","literal")';raw=workbook_bytes([record]);w=load_workbook(BytesIO(raw),data_only=False);self.assertEqual(w.active['C2'].data_type,'s');self.assertEqual(w.active['C2'].value,record['analysisObject']);w.close();self.assertEqual(read_rows('r.xlsx',raw,self.d)[0]['values']['analysisObject'],record['analysisObject'])
 def test_xlsx_cached_formula_value_is_read_without_calculation(self):
  raw=fixture([self.row]);out=BytesIO()
  with zipfile.ZipFile(BytesIO(raw)) as source,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
   for item in source.infolist():
    data=source.read(item.filename)
    if item.filename=='xl/worksheets/sheet1.xml':
     needle=b'<c r="E2" t="inlineStr"><is><t>1</t></is></c>';self.assertIn(needle,data);data=data.replace(needle,b'<c r="E2" t="str"><f>1+1</f><v>1</v></c>')
    dest.writestr(item,data)
  self.assertEqual(read_rows('cached.xlsx',out.getvalue(),self.d)[0]['values']['riskLevel'],'1')
 def test_optional_ooxml_row_and_cell_addresses(self):
  import re
  raw=fixture([self.row]);out=BytesIO()
  with zipfile.ZipFile(BytesIO(raw)) as source,zipfile.ZipFile(out,'w') as dest:
   for item in source.infolist():
    data=source.read(item.filename)
    if item.filename.endswith('sheet1.xml'):data=re.sub(rb' r="(?:[A-Z]+)?[0-9]+"',b'',data)
    dest.writestr(item,data)
  self.assertEqual(read_rows('addressless.xlsx',out.getvalue(),self.d)[0]['values']['coordinate'],'114.025,32.98')
 def test_imported_batches_survive_online_backup_restore(self):
  first=self.import_bytes(fixture([self.row,self.row]));self.s.delete(self.p,self.d,key(),first['ids'][0],1);path=Path(self.tmp.name)/'restore.sqlite'
  with closing(self.s.connect()) as source,closing(sqlite3.connect(path)) as dest:source.backup(dest)
  restored=RiskUnitStore(path)
  with closing(self.s.connect()) as source,closing(restored.connect()) as dest:
   for table in ('rp_risk_units','rp_risk_events','rp_risk_requests','rp_risk_outbox'):self.assertEqual([tuple(r) for r in source.execute('SELECT * FROM '+table+' ORDER BY rowid')],[tuple(r) for r in dest.execute('SELECT * FROM '+table+' ORDER BY rowid')])
   self.assertEqual(dest.execute('PRAGMA integrity_check').fetchone()[0],'ok');self.assertEqual(dest.execute('PRAGMA foreign_key_check').fetchall(),[])
  self.assertEqual(restored.history(self.p,self.d,first['ids'][0])[0]['payload']['importBatch']['id'],first['batchId'])
 def test_duplicate_zip_entry_and_out_of_range_dimensions(self):
  raw=fixture([self.row]);out=BytesIO()
  with zipfile.ZipFile(BytesIO(raw)) as source,zipfile.ZipFile(out,'w') as dest,warnings.catch_warnings():
   warnings.simplefilter('ignore',UserWarning)
   for item in source.infolist():dest.writestr(item,source.read(item.filename))
   dest.writestr('xl/workbook.xml',source.read('xl/workbook.xml'))
  self.error(400,read_rows,'r.xlsx',out.getvalue(),self.d)
  for old,new in [(b'r="A2"',b'r="XFD2"'),(b'<row r="2">',b'<row r="999999">')]:
   out=BytesIO()
   with zipfile.ZipFile(BytesIO(raw)) as source,zipfile.ZipFile(out,'w') as dest:
    for item in source.infolist():dest.writestr(item,source.read(item.filename).replace(old,new) if item.filename.endswith('sheet1.xml') else source.read(item.filename))
   self.error(413,read_rows,'r.xlsx',out.getvalue(),self.d)
 def test_formula_without_cache_boolean_and_extra_data_rejected(self):
  for value in ('=1+1',True):
   row=list(self.row);row[4]=value;self.error(400,self.import_bytes,fixture([row]))
  self.error(400,self.import_bytes,fixture([[*self.row,'unexpected']]))

if __name__=='__main__':unittest.main(verbosity=2)
