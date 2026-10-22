"""The reference RiskLevel six-column XLS/XLSX contract, without executing Excel."""
from io import BytesIO,StringIO
from pathlib import PurePath
import math,posixpath,re,zipfile
from services.reference_errors import DemoError

HEADERS=('部门名称','负责人','分析对象','分析单元','风险等级','坐标')
FIELDS=('department','person','analysisObject','analysisUnit','riskLevel','coordinate')
FILE_LIMIT=100*1024*1024
ROW_LIMIT=10000
XML_LIMIT=64*1024*1024
ZIP_TOTAL_LIMIT=256*1024*1024
XLSX_MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

def _cell(value,row):
 if value is None:return ''
 if isinstance(value,str):
  if len(value)>256:raise DemoError(400,f'第{row}行：单元格超过256字符')
  return value.strip()
 if type(value) in (int,float) and math.isfinite(value):return format(value,'.15g')
 raise DemoError(400,f'第{row}行：单元格需为文本或有限数字')

def _xlsx_preflight(raw):
 from defusedxml import ElementTree as ET
 with zipfile.ZipFile(BytesIO(raw)) as z:
  info=z.infolist();names=[i.filename for i in info]
  if len(names)!=len(set(names)):raise DemoError(400,'Excel压缩包存在重复文件')
  if len(info)>2000 or sum(i.file_size for i in info)>ZIP_TOTAL_LIMIT or any(i.file_size>XML_LIMIT for i in info):raise DemoError(413,'Excel展开内容超过候选容量上限')
  if any(i.flag_bits&1 for i in info):raise DemoError(400,'不支持加密的Excel文件')
  def metadata(name):
   if z.getinfo(name).file_size>2*1024*1024:raise DemoError(413,'Excel元数据过大')
   return ET.fromstring(z.read(name))
  book=metadata('xl/workbook.xml');rels=metadata('xl/_rels/workbook.xml.rels');ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
  first=book.find('s:sheets/s:sheet',ns)
  if first is None:raise DemoError(400,'Excel没有工作表')
  relationship=first.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id');matches=[r for r in rels if r.attrib.get('Id')==relationship]
  if len(matches)!=1 or matches[0].attrib.get('TargetMode')=='External' or not matches[0].attrib.get('Type','').endswith('/worksheet'):raise DemoError(400,'Excel首个页签需为普通工作表')
  target=matches[0].attrib['Target'];path=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl',target))
  if not path.startswith('xl/worksheets/') or not path.endswith('.xml'):raise DemoError(400,'Excel工作表路径无效')
  previous_row=0;current_row=None;seen_columns=set();previous_column=0
  with z.open(path) as stream:
   for event,element in ET.iterparse(stream,events=('start','end')):
    tag=element.tag.rsplit('}',1)[-1]
    if event=='start' and tag=='row':
     number=element.attrib.get('r',str(previous_row+1))
     if not re.fullmatch('[0-9]{1,7}',number):raise DemoError(400,'Excel行号无效')
     current_row=int(number);seen_columns=set();previous_column=0
     if current_row<=previous_row:raise DemoError(400,'Excel存在重复或倒序行')
     if current_row>ROW_LIMIT+1:raise DemoError(413,'单次最多导入10000行数据，请拆分文件')
     previous_row=current_row
    if event=='start' and tag=='c':
     reference=element.attrib.get('r')
     if reference is None:
      if current_row is None:raise DemoError(400,'Excel单元格位置无效')
      col=previous_column+1
     else:
      match=re.fullmatch('([A-Z]{1,3})([0-9]{1,7})',reference)
      if not match or current_row!=int(match[2]):raise DemoError(400,'Excel单元格位置无效')
      col=0
      for c in match[1]:col=col*26+ord(c)-64
     if col>64:raise DemoError(413,'Excel最多支持64列，请使用下载模板')
     if col in seen_columns:raise DemoError(400,f'第{current_row}行：存在重复单元格')
     seen_columns.add(col);previous_column=col
    if event=='end':element.clear()

def read_rows(filename,raw,directory):
 if not isinstance(filename,str) or len(filename)>200 or any(ord(c)<32 for c in filename) or PurePath(filename).name!=filename or '/' in filename or '\\' in filename:raise DemoError(400,'Excel文件名无效')
 if not isinstance(raw,bytes) or not 1<=len(raw)<FILE_LIMIT:raise DemoError(413,'Excel需小于100MiB')
 extension=PurePath(filename).suffix.lower();entries=[]
 try:
  if extension=='.xlsx':
   if not raw.startswith(b'PK\x03\x04'):raise DemoError(400,'文件内容与xlsx扩展名不符')
   _xlsx_preflight(raw)
   from openpyxl import load_workbook
   from openpyxl.xml.functions import DEFUSEDXML
   if not DEFUSEDXML:raise DemoError(503,'Excel解析依赖未完整安装')
   wb=load_workbook(BytesIO(raw),read_only=True,data_only=True,keep_links=False)
   try:
    ws=wb.worksheets[0];ws.reset_dimensions()
    table=[tuple(c.value for c in row) for row in ws.iter_rows()]
   finally:wb.close()
  elif extension=='.xls':
   if not raw.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):raise DemoError(400,'文件内容与xls扩展名不符')
   import xlrd
   wb=xlrd.open_workbook(file_contents=raw,on_demand=True,ragged_rows=True,logfile=StringIO(),ignore_workbook_corruption=False)
   try:
    ws=wb.sheet_by_index(0)
    if ws.nrows>ROW_LIMIT+1 or ws.ncols>64:raise DemoError(413,'Excel超出单次10000行/64列上限')
    table=[]
    for row in range(ws.nrows):
     cells=ws.row(row)
     if any(c.ctype not in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK,xlrd.XL_CELL_TEXT,xlrd.XL_CELL_NUMBER) for c in cells):raise DemoError(400,f'第{row+1}行：存在日期、布尔或错误单元格')
     table.append(tuple(c.value for c in cells))
   finally:wb.release_resources()
  else:raise DemoError(415,'仅支持xls和xlsx文件')
  if not table:raise DemoError(400,'Excel没有表头')
  header=[_cell(x,1) for x in table[0]]
  while header and header[-1]=='':header.pop()
  if len(header)!=len(HEADERS) or set(header)!=set(HEADERS):raise DemoError(400,'表头需包含且仅包含：'+ '、'.join(HEADERS))
  positions={field:header.index(label) for field,label in zip(FIELDS,HEADERS)}
  for number,row in enumerate(table[1:],2):
   values=[_cell(v,number) for v in row]
   if not any(values):continue
   if any(values[len(header):]):raise DemoError(400,f'第{number}行：表头以外存在数据')
   values+=['']*(len(header)-len(values));form={field:values[index] for field,index in positions.items()}
   level=form['riskLevel']
   if level not in directory.risk_levels:
    codes=[k for k,v in directory.risk_levels.items() if v==level]
    if len(codes)!=1:raise DemoError(400,f'第{number}行：风险等级无法在项目字典中唯一确定')
    form['riskLevel']=codes[0]
   entries.append(dict(rowNumber=number,values=form))
  if not entries:raise DemoError(400,'Excel没有可导入的数据行')
  return entries
 except DemoError:raise
 except ImportError:raise DemoError(503,'Excel解析依赖未安装，请联系运维')
 except Exception:raise DemoError(400,'Excel损坏、加密或格式不受支持，请使用下载模板')

def workbook_bytes(records=()):
 try:
  from openpyxl import Workbook
  from openpyxl.styles import Font,PatternFill
  from openpyxl.utils import get_column_letter
 except ImportError:raise DemoError(503,'Excel生成依赖未安装，请联系运维')
 wb=Workbook();wb.remove(wb.active);ws=wb.create_sheet('sheet');ws.append(HEADERS);ws.freeze_panes='A2'
 for index,cell in enumerate(ws[1],1):
  cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='EEEEEE');ws.column_dimensions[get_column_letter(index)].width=20
 for record in records:
  ws.append([record[field] for field in FIELDS])
  for cell in ws[ws.max_row]:cell.data_type='s';cell.number_format='@'
 ws.auto_filter.ref=ws.dimensions;target=BytesIO();wb.save(target);wb.close();return target.getvalue()
